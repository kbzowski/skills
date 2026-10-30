# /// script
# requires-python = ">=3.10"
# dependencies = ["pymupdf>=1.24", "python-docx>=1.1"]
# ///
"""Build a translated DOCX + PDF from line-based markup files (see references/markup.md).

    uv run --script build_doc.py --markup-dir WORK --src SOURCE.pdf --out "OUT_BASE" [--render DIR]

Markup files are WORK/tr_*.txt, concatenated in name order. The PDF is produced by
LibreOffice; a table of contents (@toc) gets its page numbers in a second pass.
"""
import argparse
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import pymupdf
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ALIGN = {'left': WD_ALIGN_PARAGRAPH.LEFT, 'center': WD_ALIGN_PARAGRAPH.CENTER,
         'right': WD_ALIGN_PARAGRAPH.RIGHT, 'justify': WD_ALIGN_PARAGRAPH.JUSTIFY}
TALIGN = {'left': WD_TABLE_ALIGNMENT.LEFT, 'center': WD_TABLE_ALIGNMENT.CENTER, 'right': WD_TABLE_ALIGNMENT.RIGHT}
INDENT = {'n': (1.1, 1.1), 'n2': (1.9, 0.8), 'n3': (2.7, 0.8), 'n4': (4.0, 0.6)}  # left, hanging (cm)
PI = {'pi': 1.1, 'pi2': 1.9, 'pi3': 2.7}


def kv(arg):
    """Parse key=value pairs; values may be "quoted"."""
    return {k: v.strip('"') for k, v in re.findall(r'(\w+)=("[^"]*"|\S+)', arg)}


def runs(par, text, bold=False, size=None):
    """Inline markup: **bold** __underline__ //italic// ^^small caps^^; 3+ underscores stay literal."""
    st = dict(b=bold, u=False, i=False, sc=False)
    for tok in re.split(r'(_{3,}|\*\*|__|//|\^\^)', text):
        key = {'**': 'b', '__': 'u', '//': 'i', '^^': 'sc'}.get(tok)
        if key:
            st[key] = not st[key]
        elif tok:
            r = par.add_run(tok)
            r.bold, r.underline, r.italic, r.font.small_caps = st['b'], st['u'], st['i'], st['sc']
            if size:
                r.font.size = Pt(size)


def borders(table, mode, underline_cols=()):
    el = OxmlElement('w:tblBorders')
    for e in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        b = OxmlElement(f'w:{e}')
        if mode == 'none' or (mode == 'box' and e.startswith('inside')):
            b.set(qn('w:val'), 'nil')
        else:
            for k, v in (('val', 'single'), ('sz', '12' if mode == 'box' else '4'), ('color', '000000')):
                b.set(qn(f'w:{k}'), v)
        el.append(b)
    table._tbl.tblPr.append(el)
    for row in table.rows:
        for c in underline_cols:
            tcb, b = OxmlElement('w:tcBorders'), OxmlElement('w:bottom')
            for k, v in (('val', 'single'), ('sz', '4'), ('color', '000000')):
                b.set(qn(f'w:{k}'), v)
            tcb.append(b)
            row.cells[c]._tc.get_or_add_tcPr().append(tcb)


def widths(table, cms):
    table.autofit = False
    for col, w in zip(table.columns, cms):
        col.width = Cm(w)  # gridCol: LibreOffice ignores per-cell widths without it
    for row in table.rows:
        for cell, w in zip(row.cells, cms):
            cell.width = Cm(w)


class Builder:
    def __init__(self, lines, src, workdir):
        self.lines, self.src, self.workdir = lines, src, workdir
        self.page = dict(font='Times New Roman', size='12', top='2.5', bottom='2.5', left='2.0', right='2.0',
                         width='21.0', height='29.7', ls='1.0', after='6', align='justify')
        for ln in lines:
            if ln.startswith('@page '):
                self.page.update(kv(ln[6:]))
        self.text_w = float(self.page['width']) - float(self.page['left']) - float(self.page['right'])
        # @toc TEXT [|| KEY]: KEY is what to search for when the heading differs from the TOC text
        self.toc_texts = [ln[5:].split('||')[-1].strip() for ln in lines if ln.startswith('@toc ')]

    def image(self, page, index, scale_dpi=300):
        """Render an image's area from the source page (raw extraction loses soft masks -> black backgrounds)."""
        pg = self.src[int(page) - 1]
        xref = pg.get_images()[int(index)][0]
        out = self.workdir / f'img_p{page}_{index}.png'
        pg.get_pixmap(dpi=scale_dpi, clip=pg.get_image_rects(xref)[0]).save(out)
        return out

    def para(self, where, text='', align=None, left=0.0, hang=0.0, before=0, after=None, bold=False, keep=False):
        p = where.add_paragraph()
        pf = p.paragraph_format
        pf.alignment = ALIGN[align or self.page['align']]
        pf.space_before, pf.space_after = Pt(before), Pt(float(self.page['after']) if after is None else after)
        pf.left_indent, pf.first_line_indent = Cm(left), Cm(-hang)
        pf.keep_with_next, pf.line_spacing = keep, self.ls
        if hang:
            pf.tab_stops.add_tab_stop(Cm(left))
        runs(p, text, bold)
        return p

    def build(self, toc_pages):
        doc = Document()
        st = doc.styles['Normal']
        st.font.name, st.font.size = self.page['font'], Pt(float(self.page['size']))
        st.element.rPr.rFonts.set(qn('w:eastAsia'), self.page['font'])
        sec = doc.sections[0]
        sec.page_width, sec.page_height = Cm(float(self.page['width'])), Cm(float(self.page['height']))
        for side in ('top', 'bottom', 'left', 'right'):
            setattr(sec, f'{side}_margin', Cm(float(self.page[side])))
        sec.header_distance = Cm(0.9)
        self.ls = float(self.page['ls'])
        toc_i, table, tab_opts = 0, None, {}
        for raw in self.lines:
            line = raw.rstrip()
            if not line.strip() or line.startswith('#'):
                continue
            cmd, _, arg = line.partition(' ')
            if table is not None and cmd != '@row' and cmd != '@endtable':
                raise ValueError(f'expected @row or @endtable inside @table, got: {line[:60]}')
            if cmd == '@page':
                pass
            elif cmd == '@header-image':  # first page only, e.g. a stamp or logo
                o = kv(arg)
                sec.different_first_page_header_footer = True
                hp = sec.first_page_header.paragraphs[0]
                hp.alignment = ALIGN[o.get('align', 'right')]
                hp.paragraph_format.right_indent = Cm(float(o.get('indent', 0)))
                hp.add_run().add_picture(str(self.image(o['page'], o.get('index', 0))), width=Cm(float(o['width'])))
            elif cmd == '@img':
                o = kv(arg)
                p = self.para(doc, align=o.get('align', 'center'))
                p.add_run().add_picture(str(self.image(o['page'], o.get('index', 0))), width=Cm(float(o['width'])))
            elif cmd == '@cover':  # boxed title block: @cover size=14 gap=6 | LINE | LINE ...
                opts, *titles = [s.strip() for s in arg.split('|')]
                o = kv(opts)
                for _ in range(int(o.get('gap', 6))):
                    self.para(doc, after=0)
                t = doc.add_table(rows=1, cols=1); borders(t, 'box'); widths(t, [self.text_w])
                c = t.cell(0, 0); c.paragraphs[0].paragraph_format.space_after = Pt(0)
                for i, s in enumerate(titles):
                    p = self.para(c, s, align='center', before=30 if i else 24,
                                  after=30 if i == len(titles) - 1 else 6, bold=True)
                    for r in p.runs:
                        r.font.size = Pt(float(o.get('size', 14)))
            elif cmd == '@break':
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            elif cmd == '@ls':
                self.ls = float(arg)
            elif cmd == '@gap':
                self.para(doc, after=0)
            elif cmd == '@h':
                self.para(doc, f'__{arg}__', align='left', before=18, after=12, bold=True, keep=True)
            elif cmd == '@hn':  # heading without underline
                self.para(doc, arg, align='left', before=18, after=12, bold=True, keep=True)
            elif cmd == '@hc':
                self.para(doc, f'__{arg}__', align='center', before=6, after=12, bold=True, keep=True)
            elif cmd in ('@p', '@pl', '@pc', '@pr'):
                self.para(doc, arg, align={'@pl': 'left', '@pc': 'center', '@pr': 'right'}.get(cmd))
            elif cmd[1:] in PI:
                self.para(doc, arg, left=PI[cmd[1:]])
            elif cmd[1:] in INDENT:
                lab, _, txt = arg.partition(' | ')
                left, hang = INDENT[cmd[1:]]
                self.para(doc, f'{lab}\t{txt}', left=left, hang=hang)
            elif cmd == '@bul':
                self.para(doc, f'•\t{arg}', left=1.9, hang=0.8)
            elif cmd == '@list':  # tight numbered list, no space after: @list LABEL | TEXT
                lab, _, txt = arg.partition(' | ')
                self.para(doc, f'{lab}\t{txt}', align='left', left=1.9, hang=0.8, after=0)
            elif cmd == '@toc':
                p = self.para(doc, arg.split('||')[0].strip(), align='left', after=2)
                p.paragraph_format.tab_stops.add_tab_stop(Cm(self.text_w), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
                p.add_run(f'\t{toc_pages[toc_i] if toc_pages else "0"}')
                toc_i += 1
            elif cmd == '@table':
                tab_opts = kv(arg)
                table = []
            elif cmd == '@row':
                table.append([s.strip() for s in arg.split('|')])
            elif cmd == '@endtable':
                self.table(doc, table, tab_opts)
                table = None
            else:
                raise ValueError(f'unknown directive: {line[:60]}')
        if table is not None:
            raise ValueError('@table without @endtable')
        assert toc_pages is None or toc_i == len(toc_pages), 'TOC count mismatch'
        return doc

    def table(self, doc, rows, o):
        ncols = max(len(r) for r in rows)
        w = [float(x) for x in o['widths'].split(',')] if 'widths' in o else [self.text_w / ncols] * ncols
        assert len(w) == ncols, f'@table widths has {len(w)} values for {ncols} columns'
        t = doc.add_table(rows=len(rows), cols=ncols)
        borders(t, o.get('borders', 'none'), [int(c) - 1 for c in o['underline'].split(',')] if 'underline' in o else ())
        widths(t, w)
        t.alignment = TALIGN[o.get('align', 'left')]
        size = float(o['size']) if 'size' in o else None
        for r, row in enumerate(rows):
            if o.get('cantsplit', 'yes') == 'yes':
                trp = t.rows[r]._tr.get_or_add_trPr(); cs = OxmlElement('w:cantSplit'); trp.append(cs)
            for c, s in enumerate(row):
                p = t.cell(r, c).paragraphs[0]
                p.paragraph_format.alignment = ALIGN[o.get('cellalign', 'left')]
                p.paragraph_format.space_after = Pt(float(o.get('after', 2)))
                runs(p, s, size=size)
        self.para(doc, after=0).paragraph_format.line_spacing = Pt(float(o.get('gapafter', 4)))


def find_soffice(given):
    for c in [given, shutil.which('soffice'), shutil.which('libreoffice'),
              r'C:\Program Files\LibreOffice\program\soffice.exe',
              '/Applications/LibreOffice.app/Contents/MacOS/soffice', '/usr/bin/soffice']:
        if c and Path(c).exists():
            return c
    sys.exit('LibreOffice (soffice) not found; pass --soffice PATH')


def to_pdf(docx, soffice):
    profile = Path(tempfile.mkdtemp()).as_uri()  # private profile: works even if LibreOffice is open
    subprocess.run([soffice, f'-env:UserInstallation={profile}', '--headless', '--convert-to', 'pdf',
                    '--outdir', str(docx.parent), str(docx)], check=True, capture_output=True)
    return docx.parent / (docx.stem + '.pdf')


def norm(s):
    return re.sub(r'\s+', ' ', s).strip()


def toc_pages(pdf, texts):
    """Page of each TOC entry's heading: first page whose text contains the entry's prefix outside a
    dot-leader line (TOC lines carry dot leaders, headings do not). Prefix match survives wrapped headings."""
    pages = [[norm(ln) for ln in p.get_text().splitlines()] for p in pymupdf.open(pdf)]
    out = []
    for t in texts:
        key = norm(re.sub(r'(\*\*|__|//|\^\^)', '', t))[:35]
        hit = next((i + 1 for i, lines in enumerate(pages)
                    if any(key in ln and '..' not in ln and '…' not in ln for ln in lines)), None)
        if hit is None:
            sys.exit(f'TOC entry not found as a heading in the PDF: {t!r}. If the heading is worded '
                     f'differently, write: @toc TOC TEXT || HEADING PREFIX')
        out.append(hit)
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--markup-dir', required=True, type=Path)
    ap.add_argument('--src', required=True, type=Path, help='source PDF (images are taken from it)')
    ap.add_argument('--out', required=True, help='output path WITHOUT extension')
    ap.add_argument('--soffice')
    ap.add_argument('--render', type=Path, help='also save 70-dpi PNGs of every output page here')
    a = ap.parse_args()
    files = sorted(a.markup_dir.glob('tr_*.txt'))
    if not files:
        sys.exit(f'no tr_*.txt files in {a.markup_dir}')
    lines = [ln for f in files for ln in f.read_text(encoding='utf-8').splitlines()]
    b = Builder(lines, pymupdf.open(a.src), a.markup_dir)
    out = Path(a.out)
    docx = out.parent / (out.name + '.docx')  # not with_suffix: file names may contain dots
    soffice = find_soffice(a.soffice)
    b.build(None).save(docx)
    pdf = to_pdf(docx, soffice)
    pages = None
    if b.toc_texts:
        pages = toc_pages(pdf, b.toc_texts)
        b.build(pages).save(docx)
        pdf = to_pdf(docx, soffice)
        if toc_pages(pdf, b.toc_texts) != pages:
            print('WARNING: TOC page numbers shifted after filling them in; run again', file=sys.stderr)
    d = pymupdf.open(pdf)
    if a.render:
        a.render.mkdir(parents=True, exist_ok=True)
        for old in a.render.glob('out_*.png'):
            old.unlink()
        for i, p in enumerate(d):
            p.get_pixmap(dpi=70).save(a.render / f'out_{i + 1:02d}.png')
    leftovers = re.findall(r'\*\*|\^\^|(?<!:)//|^@\w+', ''.join(p.get_text() for p in d), re.M)
    print(f'ok: {docx.name} + {pdf.name}, {d.page_count} pages, toc={pages}')
    if leftovers:
        print(f'WARNING: unparsed markup in output: {sorted(set(leftovers))}', file=sys.stderr)


if __name__ == '__main__':
    main()
